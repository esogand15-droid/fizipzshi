#!/usr/bin/env python3
"""Build INVENTORY.md + work/relay/inventory.json from whatever was extracted.

Runs on the GitHub Actions runner (which has filebin access) so the agent
sandbox can plan sessions without downloading the 1.7 GB archive.
"""
from __future__ import annotations

import json
import os
import re
import subprocess
from pathlib import Path

ROOT = Path("src/extracted")
OUT_MD = Path("INVENTORY.md")
OUT_JSON = Path("work/relay/inventory.json")

VIDEO_EXT = {".mp4", ".mkv", ".avi", ".mov", ".webm", ".m4v", ".flv", ".wmv", ".ts"}
DOC_EXT = {".pdf", ".ppt", ".pptx", ".doc", ".docx", ".md", ".txt", ".rtf"}
IMG_EXT = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".gif", ".tif", ".tiff"}


def sh(cmd: list[str], timeout: int = 120) -> str:
    try:
        p = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout)
        return (p.stdout or "") + (p.stderr or "")
    except Exception as exc:  # noqa: BLE001
        return f"ERR {exc}"


def pdf_pages(path: Path) -> int | str:
    out = sh(["pdfinfo", str(path)])
    m = re.search(r"^Pages:\s+(\d+)", out, re.M)
    if m:
        return int(m.group(1))
    try:
        import pypdfium2 as pdfium

        doc = pdfium.PdfDocument(str(path))
        return len(doc)
    except Exception:  # noqa: BLE001
        return "?"


def pdf_text_chars(path: Path) -> int | str:
    try:
        import pypdfium2 as pdfium

        doc = pdfium.PdfDocument(str(path))
        total = 0
        for i in range(len(doc)):
            page = doc[i]
            tp = page.get_textpage()
            total += len(tp.get_text_range() or "")
            tp.close()
            page.close()
        return total
    except Exception:  # noqa: BLE001
        return "?"


def pptx_slides(path: Path) -> int | str:
    try:
        from pptx import Presentation

        return len(Presentation(str(path)).slides)
    except Exception:  # noqa: BLE001
        return "?"


def media_duration(path: Path) -> str:
    out = sh(
        [
            "ffprobe", "-v", "error", "-show_entries", "format=duration",
            "-of", "default=nw=1:nk=1", str(path),
        ]
    ).strip()
    try:
        secs = float(out)
    except ValueError:
        return "?"
    return f"{int(secs // 60)}m{int(secs % 60):02d}s"


def kind(p: Path) -> str:
    ext = p.suffix.lower()
    if ext in VIDEO_EXT:
        return "video"
    if ext in DOC_EXT:
        return "doc"
    if ext in IMG_EXT:
        return "image"
    return "other"


def main() -> None:
    rows: list[dict] = []
    for dirpath, _dirnames, filenames in os.walk(ROOT):
        for fn in sorted(filenames):
            p = Path(dirpath) / fn
            rel = p.relative_to(ROOT.parent.parent).as_posix()
            try:
                size = p.stat().st_size
            except OSError:
                continue
            entry: dict = {
                "path": rel,
                "name": fn,
                "size": size,
                "size_h": f"{size / 1e6:.1f}MB",
                "kind": kind(p),
            }
            if entry["kind"] == "video":
                entry["duration"] = media_duration(p)
            elif p.suffix.lower() == ".pdf":
                entry["pages"] = pdf_pages(p)
                entry["text_chars"] = pdf_text_chars(p)
            elif p.suffix.lower() == ".pptx":
                entry["slides"] = pptx_slides(p)
            rows.append(entry)
            print(entry)

    rows.sort(key=lambda r: r["path"])
    OUT_JSON.parent.mkdir(parents=True, exist_ok=True)
    OUT_JSON.write_text(json.dumps(rows, ensure_ascii=False, indent=1), encoding="utf-8")

    counts: dict[str, int] = {}
    total_video = 0.0
    for r in rows:
        counts[r["kind"]] = counts.get(r["kind"], 0) + 1
        if r["kind"] == "video" and r.get("duration", "?").endswith("s"):
            m = re.match(r"(\d+)m(\d+)s", r["duration"])
            if m:
                total_video += int(m.group(1)) * 60 + int(m.group(2))

    lines = [
        "# INVENTORY — HumsYar medical physics sources",
        "",
        f"Total files: {len(rows)} — "
        + ", ".join(f"{k}: {v}" for k, v in sorted(counts.items())),
        f"Total video duration: {int(total_video // 3600)}h{int(total_video % 3600 // 60):02d}m",
        "",
        "| # | file | kind | size | detail |",
        "|---|------|------|------|--------|",
    ]
    for i, r in enumerate(rows, 1):
        detail = ""
        if r["kind"] == "video":
            detail = r.get("duration", "?")
        elif "pages" in r:
            detail = f"{r['pages']} pages / {r.get('text_chars')} text chars"
        elif "slides" in r:
            detail = f"{r['slides']} slides"
        lines.append(f"| {i} | `{r['path']}` | {r['kind']} | {r['size_h']} | {detail} |")
    OUT_MD.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print("wrote", OUT_MD)


if __name__ == "__main__":
    main()
